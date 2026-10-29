"""Excel 导入导出。

为什么不用 pandas：这里只需要读写几十列的小表，openpyxl 更轻，且能精细控制
单元格样式与列宽，导出的模板直接能发给同事填，不需要再教格式。
"""

import io
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlmodel import Session, select

from app.models.case import Module, TestCase
from app.models.enums import CaseStatus, CaseType, Priority

STEP_PAIRS = 3

HEADERS = [
    "用例编号",
    "用例标题",
    "模块编码",
    "优先级",
    "用例类型",
    "前置条件",
    *[f"步骤{i}" for i in range(1, STEP_PAIRS + 1)],
    *[f"预期{i}" for i in range(1, STEP_PAIRS + 1)],
    "预期结果",
    "标签",
    "状态",
    "备注",
]

SAMPLE = [
    "VOICE-001",
    "语音唤醒后 1 秒内说出导航意图，应正确解析目的地",
    "VOICE",
    "P0",
    "function",
    "车机已开机，网络正常，语音助手处于待唤醒状态",
    "说“你好小牛”唤醒语音助手",
    "唤醒成功，界面出现聆听动效",
    "说“导航到武汉纺织大学”",
    "1 秒内返回路线规划结果，目的地名称与口述一致",
    "点击开始导航",
    "进入导航页面，路线开始播报",
    "语音指令被正确解析并进入导航",
    "AI语音,冒烟",
    "active",
    "需在安静环境执行",
]

VALID_TYPES = {e.value for e in CaseType}
VALID_STATUS = {e.value for e in CaseStatus}
VALID_PRIORITY = {e.value for e in Priority}


def build_template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "用例"
    ws.append(HEADERS)
    ws.append(SAMPLE)

    head_fill = PatternFill("solid", fgColor="E8EEF7")
    for idx, _ in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=idx)
        cell.font = Font(bold=True, color="1F2D3D")
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        width = 26 if "步骤" in HEADERS[idx - 1] or "预期" in HEADERS[idx - 1] else 16
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"

    guide = wb.create_sheet("填写说明")
    for line in [
        ["字段", "说明"],
        ["用例编号", "唯一，建议 模块前缀-序号，如 NAV-012。已存在则更新该用例"],
        ["模块编码", "必须是系统中已存在的模块编码（见「模块」页）"],
        ["优先级", "P0 / P1 / P2 三档，P0 为阻塞级"],
        ["用例类型", "function / performance / compatibility / stability / interface"],
        ["步骤N / 预期N", f"最多 {STEP_PAIRS} 组，按顺序填写，留空跳过"],
        ["标签", "英文逗号分隔，如 AI语音,冒烟"],
        ["状态", "active / draft / deprecated，默认 active"],
    ]:
        guide.append(line)
    guide.column_dimensions["A"].width = 18
    guide.column_dimensions["B"].width = 70
    guide["A1"].font = guide["B1"].font = Font(bold=True)

    mods = wb.create_sheet("模块")
    mods.append(["模块编码", "模块名称"])
    for code, name in [
        ("VOICE", "AI语音"),
        ("NAVI", "导航"),
        ("POI", "POI搜索"),
        ("ROUTE", "路线规划"),
        ("CLUSTER", "仪表显示"),
        ("BT", "蓝牙电话"),
        ("MEDIA", "多媒体"),
        ("SETTING", "系统设置"),
    ]:
        mods.append([code, name])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _cell_str(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def import_cases(session: Session, data: bytes, created_by: Any = None) -> dict[str, Any]:
    wb = load_workbook(io.BytesIO(data), data_only=True)
    ws = wb["用例"] if "用例" in wb.sheetnames else wb.active

    header = [_cell_str(c.value) for c in next(ws.iter_rows(min_row=1, max_row=1))]
    col = {name: header.index(name) for name in HEADERS if name in header}
    missing = [h for h in ["用例编号", "用例标题", "模块编码"] if h not in col]
    if missing:
        raise ValueError(f"表头缺少必需列：{'、'.join(missing)}")

    modules = {m.code: m for m in session.exec(select(Module)).all()}
    existing = {c.code: c for c in session.exec(select(TestCase)).all()}

    created = updated = skipped = 0
    errors: list[dict[str, Any]] = []

    for row_no, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not any(row):
            continue
        cells = [_cell_str(v) for v in row] + [""] * (len(HEADERS) - len(row))

        def get(name: str, row_cells: list[str] = cells) -> str:
            idx = col.get(name)
            return row_cells[idx] if idx is not None else ""

        code = get("用例编号")
        title = get("用例标题")
        mod_code = get("模块编码").upper()
        if not code or not title:
            skipped += 1
            continue
        if mod_code not in modules:
            errors.append({"row": row_no, "reason": f"模块编码 {mod_code or '(空)'} 不存在"})
            continue

        priority = get("优先级").upper() or "P2"
        if priority not in VALID_PRIORITY:
            priority = "P2"
        case_type = get("用例类型").lower() or "function"
        if case_type not in VALID_TYPES:
            case_type = "function"
        status = get("状态").lower() or "active"
        if status not in VALID_STATUS:
            status = "active"

        steps = []
        for i in range(1, STEP_PAIRS + 1):
            action, expected = get(f"步骤{i}"), get(f"预期{i}")
            if action or expected:
                steps.append({"no": len(steps) + 1, "action": action, "expected": expected})

        tags = [t.strip() for t in get("标签").replace("，", ",").split(",") if t.strip()]
        payload = {
            "title": title,
            "module_id": modules[mod_code].id,
            "priority": priority,
            "case_type": case_type,
            "precondition": get("前置条件"),
            "steps": steps,
            "expected": get("预期结果"),
            "tags": tags,
            "status": status,
            "remark": get("备注"),
        }

        if code in existing:
            obj = existing[code]
            for k, v in payload.items():
                setattr(obj, k, v)
            session.add(obj)
            updated += 1
        else:
            obj = TestCase(code=code, created_by=getattr(created_by, "id", None), **payload)
            session.add(obj)
            existing[code] = obj
            created += 1

    session.commit()
    return {"created": created, "updated": updated, "skipped": skipped, "errors": errors}
